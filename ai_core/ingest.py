"""
File ingestion: turn uploaded files into text, then feed RAG.

Supported out of the box:
  - text-ish : .txt .md .log .json .html .htm .xml .tsv
  - tabular  : .csv (.xlsx if `openpyxl` installed)
  - pdf      : text PDFs via `pypdf`
  - images   : .png .jpg .jpeg .webp .gif via Gemini vision (transcribes text)

    file bytes ──► extract_text / vision ──► rag.ingest(text)

Deps (only what you use): `pip install pypdf openpyxl`
"""
import asyncio
import csv
import io
import os

from .config import Config
from .backends import get_backend

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
TEXT_EXTS = {".txt", ".md", ".log", ".json", ".html", ".htm", ".xml", ".tsv", ".yaml", ".yml"}
_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
         ".webp": "image/webp", ".gif": "image/gif", ".bmp": "image/bmp"}


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def extract_text(data: bytes, filename: str) -> str:
    """Sync extraction for non-image files. Raises for unsupported/binary types."""
    ext = _ext(filename)

    if ext in TEXT_EXTS:
        return data.decode("utf-8", errors="ignore")

    if ext == "" or ext == ".csv":
        try:
            text = data.decode("utf-8", errors="ignore")
            rows = list(csv.reader(io.StringIO(text)))
            return "\n".join(", ".join(r) for r in rows)
        except Exception:
            return data.decode("utf-8", errors="ignore")

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("pip install pypdf to ingest PDFs")
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
        text = "\n".join(pages).strip()
        if not text:
            raise RuntimeError("PDF has no extractable text (likely scanned) — send pages as images instead")
        return text

    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("pip install openpyxl to ingest .xlsx")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    out.append(", ".join(cells))
        return "\n".join(out)

    raise RuntimeError(f"unsupported file type: {ext or 'unknown'}")


async def gemini_image_to_text(data: bytes, filename: str = "image.png") -> str:
    """Transcribe/describe an image for the knowledge base (via the active backend's vision)."""
    prompt = ("Transcribe this image for a restaurant knowledge base. Extract ALL text verbatim "
              "(menu items, prices, hours, addresses), then add a one-line description.")
    return await get_backend().vision(data, filename, prompt)


def _pdf_vision_sync(data: bytes) -> str:
    """OCR fallback for scanned PDFs via Gemini (reads the PDF natively)."""
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=Config.GEMINI_API_KEY)
    resp = client.models.generate_content(
        model=Config.SEARCH_MODEL,
        contents=[types.Part.from_bytes(data=data, mime_type="application/pdf"),
                  "Transcribe ALL text from this PDF verbatim. Render any tables as CSV."],
    )
    return resp.text or ""


async def file_to_text(data: bytes, filename: str, image_extractor=None) -> str:
    ext = _ext(filename)
    if ext in IMAGE_EXTS:
        extractor = image_extractor or gemini_image_to_text
        return await extractor(data, filename)
    if ext == ".pdf":
        try:
            return extract_text(data, filename)
        except RuntimeError:
            if Config.GEMINI_API_KEY:                      # scanned PDF -> vision OCR
                return await asyncio.to_thread(_pdf_vision_sync, data)
            raise
    return extract_text(data, filename)


async def ingest_file(rag, data: bytes, filename: str, *, namespace: str,
                      source: str = None, image_extractor=None, extra: dict = None) -> int:
    """Parse a file -> text -> RAG. Returns number of chunks stored."""
    text = await file_to_text(data, filename, image_extractor=image_extractor)
    return await rag.ingest(text, namespace=namespace, source=source or filename, extra=extra)
